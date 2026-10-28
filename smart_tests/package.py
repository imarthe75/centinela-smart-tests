"""Reader for a functional test package in the methodology's format:

  * `Juego_datos_pruebas_ART-F3-06_*.xlsx`: sheet `Ejecucion` (one row per internal case CI-*) + data sheets (Unidades, Personas,
    Convenciones...). The cycle placeholder `{C}` is substituted with the run's cycle value.
  * `ART-F4-03_*.docx` (optional): UAT cases, one block per `Caso «UAT-nn» — título`.

Accepts the ZIP as delivered, a directory, the xlsx alone or `.feature` files (Gherkin, see gherkin.py) with or without the xlsx. ZIP extraction rejects absolute paths, `..` and symlinks.
`write_execution_sheet()` copies the original workbook and fills Estado / Bloqueado por / Evidencia / Hallazgo ONLY for the rows
that have a result; every other row keeps what it had. The Resumen sheet is recomputed from the Estado column.
"""
from __future__ import annotations

import html
import os
import re
import shutil
import stat
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional

import openpyxl

from centinela_core.results import CaseResult, Status

EXEC_SHEET = "Ejecucion"
# Spanish values the methodology's enumeration uses in the Estado column (ART-F3-06).
ESTADO = {Status.PASSED: "pasó", Status.FAILED: "falló", Status.BLOCKED: "bloqueado", Status.NOT_EXECUTED: "pendiente"}


class PackageError(ValueError):
    pass


@dataclass
class TestCase:
    case_id: str
    kind: str                      # "CI" | "UAT" | "BDD"
    requirement: str = ""
    origin: str = ""
    criterion: str = ""
    actor: str = ""
    data: str = ""
    steps: str = ""
    expected: str = ""
    title: str = ""
    row: Optional[int] = None      # row number in the Ejecucion sheet (CI only)

    @property
    def rq(self) -> str:
        m = re.match(r"(RQ-\d+)", self.requirement or self.criterion or "")
        return m.group(1) if m else ""


@dataclass
class TestPackage:
    root: Path
    xlsx: Optional[Path]
    uat_docx: Optional[Path]
    cycle: str
    cases: List[TestCase] = field(default_factory=list)
    data: Dict[str, List[Dict[str, str]]] = field(default_factory=dict)
    conventions: Dict[str, str] = field(default_factory=dict)

    def case(self, case_id: str) -> TestCase:
        for c in self.cases:
            if c.case_id == case_id:
                return c
        raise KeyError(case_id)

    def row(self, sheet: str, key: str, value: str) -> Dict[str, str]:
        for r in self.data.get(sheet, []):
            if r.get(key) == value:
                return r
        raise KeyError(f"{sheet}: {key}={value}")


def _safe_extract(zpath: Path, dest: Path) -> None:
    with zipfile.ZipFile(zpath) as z:
        for info in z.infolist():
            name = info.filename
            if name.startswith(("/", "\\")) or ".." in Path(name).parts:
                raise PackageError(f"ruta insegura en el ZIP: {name}")
            if stat.S_ISLNK(info.external_attr >> 16):
                raise PackageError(f"enlace simbólico en el ZIP: {name}")
        z.extractall(dest)


def _find(root: Path, pattern: str) -> Optional[Path]:
    hits = sorted(p for p in root.rglob("*") if p.is_file() and re.search(pattern, p.name, re.I) and not p.name.startswith("~$"))
    return hits[0] if hits else None


def _sub(value, cycle: str) -> str:
    return ("" if value is None else str(value)).replace("{C}", cycle).strip()


def _sheet_dicts(ws, cycle: str) -> List[Dict[str, str]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    hdr = [(_sub(h, "") or f"col{i}") for i, h in enumerate(rows[0])]
    out = []
    for r in rows[1:]:
        if not any(v not in (None, "") for v in r):
            continue
        out.append({hdr[i]: _sub(v, cycle) for i, v in enumerate(r) if i < len(hdr)})
    return out


def _docx_paragraphs(path: Path) -> List[str]:
    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")
    paras = []
    for p in re.findall(r"<w:p[ >].*?</w:p>", xml, re.S):
        text = "".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", p, re.S))
        text = html.unescape(text).strip()
        if text:
            paras.append(text)
    return paras


_UAT_LABELS = {"Criterio de aceptación:": "criterion", "Actor:": "actor", "Pasos:": "steps", "Resultado esperado:": "expected"}


def _parse_uat(path: Path, cycle: str) -> List[TestCase]:
    out: List[TestCase] = []
    cur: Optional[TestCase] = None
    pending: Optional[str] = None
    for p in _docx_paragraphs(path):
        m = re.match(r"Caso\s+«(UAT-\d+)»\s*[—-]\s*(.*)", p)
        if m:
            cur = TestCase(case_id=m.group(1), kind="UAT", title=_sub(m.group(2), cycle))
            out.append(cur)
            pending = None
            continue
        if cur is None:
            continue
        if p in _UAT_LABELS:
            pending = _UAT_LABELS[p]
            continue
        if pending:
            setattr(cur, pending, _sub(p, cycle))
            if pending == "criterion":
                cur.requirement = cur.criterion.split("·")[0].strip()
            pending = None
    return out


def load(path: str, cycle: str) -> TestPackage:
    if not re.match(r"^[A-Za-z0-9_-]{1,20}$", cycle or ""):
        raise PackageError("el ciclo debe ser alfanumérico (1 a 20 caracteres)")
    src = Path(path)
    if not src.exists():
        raise PackageError(f"no existe: {path}")
    if src.is_file() and src.suffix.lower() == ".zip":
        root = Path(tempfile.mkdtemp(prefix="smart-tests-pkg-"))
        _safe_extract(src, root)
    elif src.is_file():
        root = src.parent
    else:
        root = src
    xlsx = src if src.suffix.lower() == ".xlsx" else _find(root, r"^Juego_datos_pruebas_.*\.xlsx$")
    features = [src] if src.suffix.lower() == ".feature" else sorted(p for p in root.rglob("*.feature") if p.is_file())
    if not xlsx and not features:
        raise PackageError("el paquete no tiene el libro Juego_datos_pruebas_*.xlsx ni archivos .feature")
    uat = _find(root, r"^ART-F4-03_.*\.docx$")
    pkg = TestPackage(root=root, xlsx=Path(xlsx) if xlsx else None, uat_docx=uat, cycle=cycle)
    if xlsx:
        wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
        if EXEC_SHEET not in wb.sheetnames:
            raise PackageError(f"el libro no tiene la hoja {EXEC_SHEET}")
        for name in wb.sheetnames:
            if name != EXEC_SHEET:
                pkg.data[name] = _sheet_dicts(wb[name], cycle)
        for r in pkg.data.get("Convenciones", []):
            if r.get("Concepto"):
                pkg.conventions[r["Concepto"]] = r.get("Valor", "")
        for i, r in enumerate(wb[EXEC_SHEET].iter_rows(values_only=True), start=1):
            if i == 1 or not r or not r[0] or not str(r[0]).startswith("CI-"):
                continue
            g = lambda k: _sub(r[k] if k < len(r) else "", cycle)
            pkg.cases.append(TestCase(case_id=g(0), kind="CI", requirement=g(1), origin=g(2), criterion=g(3), actor=g(4),
                                      data=g(5), steps=g(6), expected=g(7), row=i))
    if uat:
        pkg.cases.extend(_parse_uat(uat, cycle))
    if features:
        from smart_tests import gherkin
        seen = {c.case_id for c in pkg.cases}
        for f in features:
            try:
                for c in gherkin.to_cases(f, cycle, TestCase):
                    if c.case_id in seen:          # the same case already comes from the sheet: the sheet wins, the feature adds nothing
                        continue
                    seen.add(c.case_id)
                    pkg.cases.append(c)
            except (gherkin.GherkinError, UnicodeDecodeError) as e:
                raise PackageError(f"archivo .feature inválido: {e}") from e
    if not pkg.cases:
        raise PackageError("el paquete no tiene casos")
    return pkg


def _evidence_text(r: CaseResult) -> str:
    parts = [f"{e.kind}: {e.value}" + (f" ({e.note})" if e.note else "") for e in r.evidence]
    if r.substitution:
        parts.append(f"Sustitución: {r.substitution}")
    return " | ".join(parts)[:2000]


def empty(cycle: str) -> TestPackage:
    """A package with no cases, for a run made only of generic checks (args["checks"] without a package)."""
    if not re.match(r"^[A-Za-z0-9_-]{1,20}$", cycle or ""):
        raise PackageError("el ciclo debe ser alfanumérico (1 a 20 caracteres)")
    return TestPackage(root=Path("."), xlsx=None, uat_docx=None, cycle=cycle)


def write_execution_sheet(pkg: TestPackage, results: List[CaseResult], out_path: str) -> Dict[str, int]:
    """Copy the workbook and fill only the CI rows that have a result. Returns the counts written."""
    shutil.copyfile(pkg.xlsx, out_path)
    wb = openpyxl.load_workbook(out_path)
    ws = wb[EXEC_SHEET]
    hdr = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    need = ("Estado", "Bloqueado por", "Evidencia", "Hallazgo")
    missing = [h for h in need if h not in hdr]
    if missing:
        raise PackageError(f"la hoja {EXEC_SHEET} no tiene las columnas {missing}")
    by_id = {r.case_id: r for r in results}
    written = 0
    for row in range(2, ws.max_row + 1):
        cid = ws.cell(row=row, column=1).value
        r = by_id.get(str(cid)) if cid else None
        if not r:
            continue
        ws.cell(row=row, column=hdr["Estado"]).value = ESTADO[r.status]
        ws.cell(row=row, column=hdr["Bloqueado por"]).value = r.reason if not r.status.executed else None
        ws.cell(row=row, column=hdr["Evidencia"]).value = _evidence_text(r) or None
        ws.cell(row=row, column=hdr["Hallazgo"]).value = (r.detail if r.status is Status.FAILED else None)
        written += 1
    counts = {v: 0 for v in ESTADO.values()}
    for row in range(2, ws.max_row + 1):
        if not ws.cell(row=row, column=1).value:
            continue
        v = (ws.cell(row=row, column=hdr["Estado"]).value or "pendiente").strip()
        counts[v] = counts.get(v, 0) + 1
    if "Resumen" in wb.sheetnames:
        rs = wb["Resumen"]
        for row in range(1, rs.max_row + 1):
            k = rs.cell(row=row, column=1).value
            if k in counts:
                rs.cell(row=row, column=2).value = counts[k]
            elif k == "Total":
                rs.cell(row=row, column=2).value = sum(counts.values())
            elif k and str(k).startswith("Avance"):
                rs.cell(row=row, column=2).value = sum(counts.values()) - counts.get("pendiente", 0)
    wb.save(out_path)
    return {"filas_escritas": written, **counts}
