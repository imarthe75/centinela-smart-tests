"""Package reader (methodology format). Builds a synthetic workbook so the test does not depend on any client file."""
import os
import tempfile
import zipfile

import openpyxl
import pytest

from centinela_core.results import CaseResult, Evidence, Status
from smart_tests import package


def _make_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ejecucion"
    ws.append(["Caso", "Requerimiento", "Origen", "Criterio", "Actor", "Datos", "Pasos", "Resultado esperado",
               "Fecha listo para probar", "Estado", "Bloqueado por", "Evidencia", "Hallazgo"])
    ws.append(["CI-59", "RQ-09 Registrar", "Camino principal", "RQ-09 · CA1", "Admin", "ciclo {C}", "1. x", "árbol", "", "", "", "", ""])
    ws.append(["CI-60", "RQ-09 Registrar", "RQ-09 · A1", "RQ-09 · CA2", "Admin", "código QA-{C}", "1. y", "rechazo", "", "", "", "", ""])
    u = wb.create_sheet("Unidades")
    u.append(["Id", "Nombre", "Código de prueba", "Unidad superior"])
    u.append(["UA-01", "Coordinación", "QA-CA-{C}", "UA-R01"])
    c = wb.create_sheet("Convenciones")
    c.append(["Concepto", "Valor", "Nota"])
    c.append(["Contraseña válida", "Clave.Falsa.Test1", "x"])
    r = wb.create_sheet("Resumen")
    for row in [["Estado", "Casos"], ["pasó", 0], ["falló", 0], ["bloqueado", 0], ["pendiente", 2], ["Total", 2], ["Avance (sin pendientes)", 0]]:
        r.append(row)
    wb.save(path)


@pytest.fixture
def xlsx(tmp_path):
    p = tmp_path / "Juego_datos_pruebas_ART-F3-06_DEMO.xlsx"
    _make_xlsx(str(p))
    return str(p)


def test_cycle_substitution_and_cases(xlsx):
    pkg = package.load(xlsx, "c9")
    assert len(pkg.cases) == 2
    assert pkg.case("CI-59").data == "ciclo c9"
    assert pkg.case("CI-60").data == "código QA-c9"
    assert pkg.case("CI-59").rq == "RQ-09"
    assert pkg.row("Unidades", "Id", "UA-01")["Código de prueba"] == "QA-CA-c9"
    assert pkg.conventions["Contraseña válida"] == "Clave.Falsa.Test1"


def test_cycle_must_be_alphanumeric(xlsx):
    with pytest.raises(package.PackageError):
        package.load(xlsx, "c/9")


def test_write_fills_only_cases_with_results(xlsx, tmp_path):
    pkg = package.load(xlsx, "c1")
    results = [
        CaseResult("CI-59", Status.PASSED, evidence=[Evidence("http", "GET /x -> 200")]),
        CaseResult("CI-60", Status.BLOCKED, reason="sin datos"),
    ]
    out = str(tmp_path / "out.xlsx")
    counts = package.write_execution_sheet(pkg, results, out)
    assert counts == {"filas_escritas": 2, "pasó": 1, "falló": 0, "bloqueado": 1, "pendiente": 0}
    wb = openpyxl.load_workbook(out)
    ws = wb["Ejecucion"]
    assert ws.cell(row=2, column=10).value == "pasó"
    assert ws.cell(row=3, column=10).value == "bloqueado" and ws.cell(row=3, column=11).value == "sin datos"
    rs = {wb["Resumen"].cell(row=i, column=1).value: wb["Resumen"].cell(row=i, column=2).value for i in range(1, 8)}
    assert rs["pasó"] == 1 and rs["pendiente"] == 0 and rs["Total"] == 2


def test_zip_rejects_path_traversal(tmp_path):
    z = tmp_path / "bad.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("../evil.txt", "x")
    with pytest.raises(package.PackageError):
        package.load(str(z), "c1")


def test_missing_execution_sheet(tmp_path):
    p = tmp_path / "Juego_datos_pruebas_x.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Otra"
    wb.save(str(p))
    with pytest.raises(package.PackageError):
        package.load(str(p), "c1")
